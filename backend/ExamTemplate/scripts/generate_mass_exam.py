#!/usr/bin/python

import argparse
import re
import subprocess
import os
import sys
import shutil
import csv
try:
    import openpyxl
    OPENPYXL=True
except Exception as e:
    OPENPYXL=False
    print("Unable to import openpyxl, falling back to CSV.")

ALL_LANGUAGES= ["de", "en"]
OUT_MAPPING_NAME = "anwesenheitsliste.csv"

class Job(object):
    def __init__(self, languages, blankexamscount, examList):
        super(Job, self).__init__()
        self.languages = languages
        self.blankexamscount = blankexamscount
        self.examList = examList

    def getLanguageNumberPrefix(self, language):
        return ALL_LANGUAGES.index(language)+1

def adaptMeta(language, randomnumber, loesung, name, matrk, seq, targetPath):
    os.chdir(targetPath)
    with open('meta-exam.tex', 'r') as file:
        data = file.read()

    # viktor: alle leerzeichen escpaen
    name = re.sub(" ", "\\ ", name)

    data = re.sub("\\\\newcommand{\\\\randomexamnumber}{[0-9A-Z]+}", r'\\newcommand{\\randomexamnumber}{'+randomnumber+'}', data)
    data = re.sub("\\\\newcommand{\\\\sprache}{[den]+}", r"\\newcommand{\\sprache}{"+language+"}", data)
    data = re.sub("\\\\newcommand{\\\\zeigeloesung}{[yesno]+}", r"\\newcommand{\\zeigeloesung}{"+loesung+"}", data)
    data = re.sub("\\\\newcommand{\\\\sequenznummer}{[0-9]*}", r"\\newcommand{\\sequenznummer}{"+str(seq)+"}", data)
    data = re.sub("\\\\newcommand{\\\\vollername}{[^}]*}", r"\\newcommand{\\vollername}{"+name+"}", data)
    data = re.sub("\\\\newcommand{\\\\matrikelnummer}{[^}]*}", r"\\newcommand{\\matrikelnummer}{"+str(matrk)+"}", data)

    with open('meta-exam.tex', 'w') as file:
        file.writelines( data )

def build(targetPath, it=3, src="exam.tex"):
#    subprocess.call(["touch", "exam.pdf"], stdout=open(os.devnull, 'wb'))
#    return
    os.chdir(targetPath)
    for i in range(it):
        if subprocess.call(["pdflatex", "-halt-on-error", src], stdout=open(os.devnull, 'wb')) != 0:
            raise Exception("pdflatex returned non-zero exit code")

def generateSolution(job, targetPath, outDir):
    print(" Generating solution...")
    assert len(job.languages) >= 1
    adaptMeta(job.languages[0], "4242", "yes", "", "", "", targetPath) # uses first language in list (de if existing, en otherwise)
    build(targetPath)
    shutil.copy(os.path.join(targetPath, "exam.log"), os.path.join(outDir, "exam_solution.log"))
    shutil.copy(os.path.join(targetPath, "exam.pdf"), os.path.join(outDir, "exam_solution.pdf"))
    print(" Generated exam_solution.pdf and exam_solution.log")

def base36encode(number, alphabet='0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'):
    """Converts an integer to a base36 string."""
    if not isinstance(number, (int)):
        raise TypeError('number must be an integer')   

    if number < 0:
        raise TypeError('number must be positive')

    base36 = ''

    if 0 <= number < len(alphabet):
        return alphabet[number]

    while number != 0:
        number, i = divmod(number, len(alphabet))
        base36 = alphabet[i] + base36

    return base36


def checksum(number, alphabet='0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'):
    """Calculate the checksum. A valid number should have a checksum of 1."""
    modulus = len(alphabet)
    check = modulus // 2
    for n in number:
        check = (((check or modulus) * 2) % (modulus + 1) + alphabet.index(n)) % modulus
    return check

def calc_check_digit(number, alphabet='0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'):
    """With the provided number, calculate the extra digit that should be
    appended to make it a valid number."""
    modulus = len(alphabet)
    return alphabet[(1 - ((checksum(number, alphabet) or modulus) * 2) % (modulus + 1)) % modulus]

def genRandomNumber(languagePrefix, counter):
    while len(counter) < 4:
        counter = "0"+counter
    num = base36encode(int(languagePrefix+counter))
    parity = calc_check_digit(num);
    return num+parity

class OutputEntry(object):
    def __init__(self, lang, randomnumber, sequenznummer=None, matrikelnummer=None, name=None):
        super(OutputEntry, self).__init__()
        self.lang = lang
        self.randomnumber = randomnumber
        self.sequenznummer = sequenznummer
        self.matrikelnummer = matrikelnummer
        self.name = name

def generateOutput(outputMapping, outDir):
    sequencedDict = {}
    unassignedExams = []

    for entry in outputMapping:
        if entry.matrikelnummer == None:
            unassignedExams.append(entry)
            continue
        if entry.sequenznummer in sequencedDict:
            assignedEntry = sequencedDict[entry.sequenznummer]
            assert assignedEntry.matrikelnummer == entry.matrikelnummer
            assert assignedEntry.sequenznummer == entry.sequenznummer
            assert assignedEntry.name == entry.name
            assignedEntry.randomnumber.append(entry.randomnumber)
            continue
        entry.randomnumber = [entry.randomnumber]
        sequencedDict[entry.sequenznummer] = entry
    with open(os.path.join(outDir, OUT_MAPPING_NAME), "w") as out:
        writer = csv.writer(out, quoting=csv.QUOTE_ALL)
        writer.writerow(['Sitzplatz', 'Random', 'Matrikelnr', 'Name', 'Anwesend? (X)'])
        for key in sorted(sequencedDict):
            entry = sequencedDict[key]
            randomnumbers = "/".join(entry.randomnumber)
            writer.writerow([key, randomnumbers, entry.matrikelnummer, entry.name, ""])
        for v in unassignedExams:
            writer.writerow(['', v.randomnumber])
        #out.write(serialized.encode('utf8'))
    print("\n Wrote mapping to {}. KEEP THAT FILE!".format(OUT_MAPPING_NAME))

def generatePdf(language, job, outputMapping, targetPath, outDir, skipPdf=False, tearoffsheet=False):
    print(" Generating language "+language)
    tempfiles = []
    totalNum = len(job.examList)+job.blankexamscount
    for i in range(1, totalNum+1):
        print("  Generating exam {0} of {1} ({2:.2%})".format(i, totalNum, i/float(totalNum)))
        randomnumber = genRandomNumber(str(job.getLanguageNumberPrefix(language)), str(i))
        if i <= len(job.examList):
            examEntry = job.examList[i-1]
            name = "{} {}".format(examEntry[2], examEntry[1])
            matrikelnummer = examEntry[0]
            if not(skipPdf):
                adaptMeta(language, randomnumber, "no", name, matrikelnummer, i, targetPath)
            outputMapping.append(OutputEntry(language, randomnumber, i, matrikelnummer, name))
        else:
            if not(skipPdf):
                adaptMeta(language, randomnumber, "no", "", "", "", targetPath)
            outputMapping.append(OutputEntry(language, randomnumber))
        if not(skipPdf):
            if tearoffsheet:
                build(targetPath, it=1, src="tear-off-sheet.tex")
            build(targetPath)

        outfilename = "exam_{0}_{1}.pdf".format(language, randomnumber)
        if not(skipPdf):
            if tearoffsheet:
                tearoffname = "tearoff_{0}_{1}.pdf".format(language, randomnumber)
                shutil.copy(os.path.join(targetPath, "tear-off-sheet.pdf"), os.path.join(outDir, tearoffname))
                tempfiles.append(tearoffname)
            shutil.copy(os.path.join(targetPath, "exam.pdf"), os.path.join(outDir, outfilename))
        tempfiles.append(outfilename)

    examlogFileName = "exam_{0}.log".format(language)
    if not(skipPdf):
        shutil.copy(os.path.join(targetPath, "exam.log"), os.path.join(outDir, examlogFileName))
    print(" Copied {0}".format(examlogFileName))

    os.chdir(outDir)
    mergeArguments = ["gs", "-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite", "-q", "-sOutputFile=exam_{0}.pdf".format(language)]+tempfiles
    print(" Merging exam_{0}.pdf".format(language))
    os
    if skipPdf:
        return
    if subprocess.call(mergeArguments) != 0:
        raise Exception("Merging failed with errorcode")
    for tmpfilename in tempfiles:
        os.remove(tmpfilename)

def parseCsvStudentList(file):
    csvreader = csv.reader(file)
    while next(csvreader)[0] !="startHISsheet":
        pass
    headerColumn = next(csvreader)
    expectedHeaders = ("Matrikelnummer", "Nachname", "Vorname")

    for i, h in enumerate(expectedHeaders):
        if h not in headerColumn:
            raise Exception("Examlist column {} is {} instead {}".format(i, headerColumn[i], h))
    allStudents = []

    #cleancsv =  "\n".join(",".join(e) for e in list(csvreader))
    #studentlist = csv.DictReader(cleancsv.splitlines(), delimiter=',')
    #print(cleancsv)
    for row in csvreader:
        if row[0] == "endHISsheet":
            break
        try:
            #convert matrikl to int
            allStudents.append([int(row[5]), row[3], row[4]])
        except Exception as e:
            print(e)
            continue

    return allStudents


def parseStudentList(fileName):
    wb = openpyxl.load_workbook(fileName)
    sheet = wb[wb.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)

    while next(rows)[0] !="startHISsheet":
        pass
    headerColumn = next(rows)
    if not (headerColumn[5] == "Matrikelnummer" and headerColumn[3] == "Nachname" and headerColumn[4] == "Vorname"):
        raise Exception("Examlist columns do not match. Expected 'Nachname' at 4, 'Vorname' at 5 and 'Matrikelnummer' at 6 (numbered from 1).")
    allStudents = []

    for row in rows:
        try:
            #convert matrikl to int
            allStudents.append([int(str(row[5]).strip('\'\" ,.')), row[3], row[4]])
        except Exception as e:
            break

    return allStudents

def main():
    parser = argparse.ArgumentParser(description='Generate pdfs consisting of all exams for the given languages.')
    parser.add_argument('--noPdf', action='store_true', default=False,
                       help='Skip generating PDFs and logs, only generate the attendance list.')
    parser.add_argument('--de', action='store_true',
                       help='Generate German exam.')
    parser.add_argument('--en', action='store_true',
                       help='Generate English exam.')
    parser.add_argument('--tearoffsheet', action='store_true',
                        help='Generate tear off sheets with codes.')
    parser.add_argument('--examlist',
                       help='Personalize with exam list.')
    parser.add_argument('--blankexamscount', type=int, default=0,
                       help='Number of exams without name.')
    parser.add_argument('--outdir', default="./",
                       help='Output path.')
    parser.add_argument('--examdir', default="./",
                       help='TeX source path.')
    args = parser.parse_args()
#    if args.gennum:
#        numbers = []
#        for i in range(1000):
#            numbers.append(genRandomNumber("1", str(i)))
#            numbers.append(genRandomNumber("2", str(i)))
#        print(numbers)
#        raise SystemExit()

    if args.noPdf:
        print("Skipping PDF generation")

    if args.tearoffsheet:
        print("Generating tear-off sheets")

    if not(args.de) and not(args.en):
        print("Please give at least one language.")
        return
    if args.examlist == None and (args.blankexamscount < 1 or args.blankexamscount > 500):
        print("Do you really want to print {0} exams?!".format(args.blankexamscount))
        return
    languages = []
    if args.de:
        languages.append("de")
    if args.en:
        languages.append("en")
    examList = []
    if args.examlist:
        if os.path.exists(OUT_MAPPING_NAME):
            print("File {} already exists. Won't overwrite.".format(OUT_MAPPING_NAME))
            raise SystemExit()

        if args.examlist.endswith(".xlsx") and OPENPYXL:
                examList = parseStudentList(args.examlist)
        elif args.examlist.endswith(".csv"):
            try:
                examlistFile = open(args.examlist, "r")
                examList = parseCsvStudentList(examlistFile)
            except UnicodeDecodeError as e:
                examlistFile = open(args.examlist, mode="r", encoding="Windows-1252")
                examList = parseCsvStudentList(examlistFile)
        else:
            print("Without openpyxl I can only read csv exam list files!")
            sys.exit(1)

        # sort list!
        examList.sort(key=lambda x:x[0])
        # convert columns back to strings
        examList = [[str(x) for x in entry] for entry in examList]
        print("Parsed list with {} students".format(len(examList)))
    job = Job(languages, args.blankexamscount, examList)
    outputMapping = []
    outdir = os.path.realpath(args.outdir)
    target = os.path.realpath(args.examdir)

    for language in languages:
        generatePdf(language, job, outputMapping, target, outdir, args.noPdf, args.tearoffsheet)
    if not(args.noPdf):
        generateSolution(job, target, outdir)
    if args.examlist:
        generateOutput(outputMapping, outdir)

if __name__ == '__main__':
    main()
